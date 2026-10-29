import os
import sys
import glob
import re
import docx
import fitz

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"C:\Users\bhara\Downloads\DIGITAL JOURNALISM WEBSITE-20261022T013821Z-1-001\DIGITAL JOURNALISM WEBSITE\Blog"

print("=== BLOG / ARADHANA ===")
doc = docx.Document(os.path.join(base_dir, "Aradhana.docx"))
aradhana_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
print("Original paragraphs count:", len([p.text for p in doc.paragraphs if p.text.strip()]))
print(aradhana_text[:500])
print("...")
print(aradhana_text[-500:])

print("\n=== BLOG / DHANUSHREE ===")
pdf = fitz.open(os.path.join(base_dir, "Dhanushree.pdf"))
dh_text = ""
for page in pdf:
    dh_text += page.get_text() + "\n"
print("Original length:", len(dh_text))
print(dh_text[:500])
print("...")
print(dh_text[-500:])

print("\n=== BLOG / MITHRAA ===")
pdf_m = fitz.open(os.path.join(base_dir, "Mithraa.pdf"))
m_text = ""
for page in pdf_m:
    m_text += page.get_text() + "\n"
print("Original length:", len(m_text))
print(m_text[:500])
print("...")
print(m_text[-500:])
